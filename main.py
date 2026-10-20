import tkinter as tk
from tkinter import messagebox
import json 
from openpyxl import Workbook
import smtplib
from email.message import EmailMessage
import shutil
import os

CORREO_EMISOR = os.getenv("CORREO_EMISOR")
CONTRASENA_APP = os.getenv("CONTRASENA_APP")
lista_productos = []
def guardar_productos():
    with open("productos.json", "w", encoding="utf-8") as archivo:
        json.dump(lista_productos, archivo, ensure_ascii=False, indent=4)

def cargar_productos():
    try:
        with open("productos.json", "r", encoding="utf-8") as archivo:
            datos = json.load(archivo)
            lista_productos.extend(datos)
    except FileNotFoundError:
        pass


cargar_productos()

lista_clientes = []

def guardar_clientes():
    with open("clientes.json", "w", encoding="utf-8") as archivo:
        json.dump(lista_clientes, archivo, ensure_ascii=False, indent=4)

def exportar_clientes_excel():
    archivo = Workbook()
    hoja = archivo.active
    hoja.title = "Clientes"

    hoja.append(["Nombre", "NIT", "Dirección"])

    for cliente in lista_clientes:
        hoja.append([
            cliente["nombre"],
            cliente["nit"],
            cliente["direccion"]
        ])

    archivo.save("clientes.xlsx")

    messagebox.showinfo(
        "Excel",
        "Clientes exportados correctamente a clientes.xlsx"
    )

def cargar_clientes():
    try:
        with open("clientes.json", "r", encoding="utf-8") as archivo:
            datos = json.load(archivo)
            lista_clientes.extend(datos)
    except FileNotFoundError:
        pass 

cargar_clientes()

lista_pedidos = []

def guardar_pedidos():
    with open("pedidos.json", "w", encoding="utf-8") as archivo:
        json.dump(lista_pedidos, archivo, ensure_ascii=False, indent=4)

def cargar_pedidos():
    try:
        with open("pedidos.json", "r", encoding="utf-8") as archivo:
            datos = json.load(archivo)
            lista_pedidos.extend(datos)
    except FileNotFoundError:
        pass

cargar_pedidos()
def actualizar_excel():
    archivo = Workbook()

    hoja_clientes = archivo.active
    hoja_clientes.title = "Clientes"

    hoja_clientes.append(["Nombre", "NIT", "Dirección"])

    for cliente in lista_clientes:
        hoja_clientes.append([
            cliente["nombre"],
            cliente["nit"],
            cliente["direccion"]
        ])

    hoja_productos = archivo.create_sheet("Productos")

    hoja_productos.append(["Producto", "Precio", "Existencia"])

    for producto in lista_productos:
        hoja_productos.append([
            producto["nombre"],
            producto["precio"],
            producto["existencia"]
        ])

    hoja_ventas = archivo.create_sheet("Ventas")

    hoja_ventas.append([
        "Cliente",
        "Producto",
        "Cantidad",
        "Total"
    ])

    for pedido in lista_pedidos:
        hoja_ventas.append([
            pedido["cliente"],
            pedido["producto"],
            pedido["cantidad"],
            pedido["total"]
        ])

    hoja_resumen = archivo.create_sheet("Resumen")

    hoja_resumen.append(["INFORME GENERAL"])
    hoja_resumen.append(["Clientes registrados", len(lista_clientes)])
    hoja_resumen.append(["Productos registrados", len(lista_productos)])
    hoja_resumen.append(["Ventas realizadas", len(lista_pedidos)])

    ventas_totales = 0

    for pedido in lista_pedidos:
        ventas_totales = ventas_totales + float(pedido["total"])

    hoja_resumen.append(["Ventas totales", ventas_totales])

    archivo.save("sistema_ventas.xlsx")



def agregar_producto():
    nombre = entrada_nombre.get()
    precio = entrada_precio.get()
    existencia = entrada_existencia.get()

    if nombre == "" or precio == "" or existencia == "":
        messagebox.showwarning(
            "Advertencia",
            "Debe completar todos los campos"
        )
        return

    producto = {
        "nombre": nombre,
        "precio": precio,
        "existencia": existencia
    }

    lista_productos.append(producto)
    guardar_productos()
    actualizar_excel()
    messagebox.showinfo(
        "Correcto",
        "Producto guardado correctamente"
    )

    entrada_nombre.delete(0, tk.END)
    entrada_precio.delete(0, tk.END)
    entrada_existencia.delete(0, tk.END)




def listar_productos():

    ventana_lista = tk.Toplevel()
    ventana_lista.title("Lista de productos")
    ventana_lista.geometry("500x400")

    titulo = tk.Label(
        ventana_lista,
        text="LISTA DE PRODUCTOS",
        font=("Arial", 18)
    )

    titulo.pack(pady=15)

    if len(lista_productos) == 0:

        mensaje = tk.Label(
            ventana_lista,
            text="No hay productos registrados",
            font=("Arial", 12)
        )

        mensaje.pack(pady=20)

    else:

        for producto in lista_productos:

            texto = (
                "Producto: " + producto["nombre"] +
                "\nPrecio: Q" + str(producto["precio"]) +
                "\nExistencia: " + str(producto["existencia"]) +
                "\n-----------------------------"
            )

            etiqueta = tk.Label(
                ventana_lista,
                text=texto,
                font=("Arial", 11),
                justify="left"
            )

            etiqueta.pack(pady=5)


def editar_producto():

    if len(lista_productos) == 0:
        messagebox.showwarning(
            "Advertencia",
            "No hay productos para editar"
        )
        return

    ventana_editar = tk.Toplevel()
    ventana_editar.title("Editar producto")
    ventana_editar.geometry("400x350")

    tk.Label(
        ventana_editar,
        text="Nombre del producto:"
    ).pack(pady=5)

    entrada_editar_nombre = tk.Entry(
        ventana_editar,
        width=30
    )

    entrada_editar_nombre.pack()

    tk.Label(
        ventana_editar,
        text="Nuevo precio:"
    ).pack(pady=5)

    entrada_editar_precio = tk.Entry(
        ventana_editar,
        width=30
    )

    entrada_editar_precio.pack()

    tk.Label(
        ventana_editar,
        text="Nueva existencia:"
    ).pack(pady=5)

    entrada_editar_existencia = tk.Entry(
        ventana_editar,
        width=30
    )

    entrada_editar_existencia.pack()

    def guardar_cambios():

        nombre = entrada_editar_nombre.get()

        encontrado = False

        for producto in lista_productos:

            if producto["nombre"] == nombre:

                producto["precio"] = entrada_editar_precio.get()
                producto["existencia"] = entrada_editar_existencia.get()
                guardar_productos()
                actualizar_excel()
                encontrado = True

                messagebox.showinfo(
                    "Correcto",
                    "Producto actualizado correctamente"
                )

                ventana_editar.destroy()

                break

        if encontrado == False:

            messagebox.showerror(
                "Error",
                "No se encontró el producto"
            )

    tk.Button(
        ventana_editar,
        text="Guardar cambios",
        command=guardar_cambios,
        width=25
    ).pack(pady=20)



def eliminar_producto():

    if len(lista_productos) == 0:
        messagebox.showwarning(
            "Advertencia",
            "No hay productos para eliminar"
        )
        return

    ventana_eliminar = tk.Toplevel()
    ventana_eliminar.title("Eliminar producto")
    ventana_eliminar.geometry("400x250")

    tk.Label(
        ventana_eliminar,
        text="Nombre del producto que desea eliminar:"
    ).pack(pady=15)

    entrada_eliminar = tk.Entry(
        ventana_eliminar,
        width=30
    )

    entrada_eliminar.pack()

    def eliminar():

        nombre = entrada_eliminar.get()

        encontrado = False

        for producto in lista_productos:

            if producto["nombre"] == nombre:

                lista_productos.remove(producto)
                guardar_productos()
                actualizar_excel()
                encontrado = True

                messagebox.showinfo(
                    "Correcto",
                    "Producto eliminado correctamente"
                )

                ventana_eliminar.destroy()

                break

        if encontrado == False:

            messagebox.showerror(
                "Error",
                "No se encontró el producto"
            )

    tk.Button(
        ventana_eliminar,
        text="Eliminar producto",
        command=eliminar,
        width=25
    ).pack(pady=20)



def productos():

    ventana_productos = tk.Toplevel()

    ventana_productos.title("Productos")

    ventana_productos.geometry("500x650")

    titulo = tk.Label(
        ventana_productos,
        text="GESTIÓN DE PRODUCTOS",
        font=("Arial", 20)
    )

    titulo.pack(pady=20)


    tk.Label(
        ventana_productos,
        text="Nombre del producto:"
    ).pack(pady=5)

    global entrada_nombre

    entrada_nombre = tk.Entry(
        ventana_productos,
        width=35
    )

    entrada_nombre.pack()


    tk.Label(
        ventana_productos,
        text="Precio:"
    ).pack(pady=5)

    global entrada_precio

    entrada_precio = tk.Entry(
        ventana_productos,
        width=35
    )

    entrada_precio.pack()


    tk.Label(
        ventana_productos,
        text="Existencia:"
    ).pack(pady=5)

    global entrada_existencia

    entrada_existencia = tk.Entry(
        ventana_productos,
        width=35
    )

    entrada_existencia.pack()


    boton_agregar = tk.Button(
        ventana_productos,
        text="Agregar producto",
        command=agregar_producto,
        width=25
    )

    boton_agregar.pack(pady=10)

    boton_editar = tk.Button(
        ventana_productos,
        text="Editar producto",
        command=editar_producto,
        width=25
    )

    boton_editar.pack(pady=10)


    boton_eliminar = tk.Button(
        ventana_productos,
        text="Eliminar producto",
        command=eliminar_producto,
        width=25
    )

    boton_eliminar.pack(pady=10)

    boton_listar = tk.Button(
        ventana_productos,
        text="Listar productos",
        command=listar_productos,
        width=25
    )

    boton_listar.pack(pady=10)

    boton_cotizacion = tk.Button(
        ventana_productos,
        text="Crear cotización",
        command=mostrar_cotizacion,
        width=25
    )
    boton_cotizacion.pack(pady=10)


def listar_clientes():

    ventana_lista = tk.Toplevel()

    ventana_lista.title("Lista de clientes")

    ventana_lista.geometry("500x400")

    titulo = tk.Label(
        ventana_lista,
        text="LISTA DE CLIENTES",
        font=("Arial", 18)
    )

    titulo.pack(pady=15)

    if len(lista_clientes) == 0:

        mensaje = tk.Label(
            ventana_lista,
            text="No hay clientes registrados",
            font=("Arial", 12)
        )

        mensaje.pack(pady=20)

    else:

        for cliente in lista_clientes:

            texto = (
                "Cliente: " + cliente["nombre"] +
                "\nNIT: " + cliente["nit"] +
                "\nDirección: " + cliente["direccion"] +
                "\n-----------------------------"
            )

            etiqueta = tk.Label(
                ventana_lista,
                text=texto,
                font=("Arial", 11),
                justify="left"
            )

            etiqueta.pack(pady=5)

def editar_cliente():

    if len(lista_clientes) == 0:

        messagebox.showwarning(
            "Advertencia",
            "No hay clientes para editar"
        )
        return
    ventana_editar = tk.Toplevel()
    ventana_editar.title("Editar cliente")
    ventana_editar.geometry("450x400")


    tk.Label(
        ventana_editar,
        text="Nombre del cliente:"
    ).pack(pady=5)


    entrada_nombre_editar = tk.Entry(
        ventana_editar,
        width=35
    )

    entrada_nombre_editar.pack()


    tk.Label(
        ventana_editar,
        text="Nuevo NIT:"
    ).pack(pady=5)


    entrada_nit_editar = tk.Entry(
        ventana_editar,
        width=35
    )

    entrada_nit_editar.pack()


    tk.Label(
        ventana_editar,
        text="Nueva dirección:"
    ).pack(pady=5)


    entrada_direccion_editar = tk.Entry(
        ventana_editar,
        width=35
    )

    entrada_direccion_editar.pack()


    def guardar_cambios():

        nombre = entrada_nombre_editar.get()
        encontrado = False
        for cliente in lista_clientes:

            if cliente["nombre"] == nombre:

                cliente["nit"] = entrada_nit_editar.get()

                cliente["direccion"] = entrada_direccion_editar.get()
                guardar_clientes()
                actualizar_excel()
                encontrado = True


                messagebox.showinfo(
                    "Correcto",
                    "Cliente actualizado correctamente"
                )


                ventana_editar.destroy()

                break


        if encontrado == False:

            messagebox.showerror(
                "Error",
                "No se encontró el cliente"
            )


    boton_guardar = tk.Button(
        ventana_editar,
        text="Guardar cambios",
        command=guardar_cambios,
        width=25
    )

    boton_guardar.pack(pady=20)

def eliminar_cliente():

    if len(lista_clientes) == 0:

        messagebox.showwarning(
            "Advertencia",
            "No hay clientes para eliminar"
        )

        return


    ventana_eliminar = tk.Toplevel()

    ventana_eliminar.title("Eliminar cliente")

    ventana_eliminar.geometry("450x300")


    tk.Label(
        ventana_eliminar,
        text="Nombre del cliente que desea eliminar:"
    ).pack(pady=15)


    entrada_eliminar = tk.Entry(
        ventana_eliminar,
        width=35
    )

    entrada_eliminar.pack()


    def eliminar():

        nombre = entrada_eliminar.get()

        encontrado = False


        for cliente in lista_clientes:

            if cliente["nombre"] == nombre:

                lista_clientes.remove(cliente)
                guardar_clientes()
                actualizar_excel()
                encontrado = True


                messagebox.showinfo(
                    "Correcto",
                    "Cliente eliminado correctamente"
                )


                ventana_eliminar.destroy()

                break


        if encontrado == False:

            messagebox.showerror(
                "Error",
                "No se encontró el cliente"
            )


    boton_eliminar = tk.Button(
        ventana_eliminar,
        text="Eliminar cliente",
        command=eliminar,
        width=25
    )

    boton_eliminar.pack(pady=20)

def clientes():

    ventana_clientes = tk.Toplevel()

    ventana_clientes.title("Clientes")

    ventana_clientes.geometry("500x550")

    titulo = tk.Label(
        ventana_clientes,
        text="GESTIÓN DE CLIENTES",
        font=("Arial", 20)
    )

    titulo.pack(pady=20)

    tk.Label(
        ventana_clientes,
        text="Nombre del cliente:"
    ).pack(pady=5)

    entrada_cliente = tk.Entry(
        ventana_clientes,
        width=35
    )

    entrada_cliente.pack()

    tk.Label(
        ventana_clientes,
        text="NIT:"
    ).pack(pady=5)

    entrada_nit = tk.Entry(
        ventana_clientes,
        width=35
    )

    entrada_nit.pack()

    tk.Label(
        ventana_clientes,
        text="Dirección:"
    ).pack(pady=5)

    entrada_direccion = tk.Entry(
        ventana_clientes,
        width=35
    )

    entrada_direccion.pack()

    def guardar_cliente():

        nombre = entrada_cliente.get()
        nit = entrada_nit.get()
        direccion = entrada_direccion.get()

        if nombre == "" or nit == "" or direccion == "":

            messagebox.showwarning(
                "Advertencia",
                "Debe completar todos los campos"
            )

            return


        cliente = {
            "nombre": nombre,
            "nit": nit,
            "direccion": direccion
        }


        lista_clientes.append(cliente)
        guardar_clientes()
        actualizar_excel()
        messagebox.showinfo(
            "Correcto",
            "Cliente guardado correctamente"
        )


        entrada_cliente.delete(0, tk.END)
        entrada_nit.delete(0, tk.END)
        entrada_direccion.delete(0, tk.END)

    boton_guardar = tk.Button(
        ventana_clientes,
        text="Guardar cliente",
        command=guardar_cliente,
        width=25
    )

    boton_guardar.pack(pady=20)

    boton_listar = tk.Button(
        ventana_clientes,
        text="Listar clientes",
        command=listar_clientes,
        width=25
    )

    boton_listar.pack(pady=10)

    boton_editar = tk.Button(
        ventana_clientes,
        text="Editar cliente",
        command=editar_cliente,
        width=25
    )

    boton_editar.pack(pady=10)

    boton_eliminar = tk.Button(
        ventana_clientes,
        text="Eliminar cliente",
        command=eliminar_cliente,
        width=25
    )

    boton_eliminar.pack(pady=10)

    boton_excel = tk.Button(
        ventana_clientes,
        text="Exportar a Excel",
        command=exportar_clientes_excel,
        width=25
    )
    boton_excel.pack(pady=10)

def buscar_cliente(nombre):

    for cliente in lista_clientes:

        if cliente["nombre"] == nombre:

            return cliente

    return None

def buscar_producto(nombre):
    for producto in lista_productos:
        if producto["nombre"] == nombre:
            return producto
    return None

def eliminar_pedido():

    if len(lista_pedidos) == 0:
        messagebox.showwarning(
            "Advertencia",
            "No hay pedidos para eliminar"
        )
        return

    ventana_eliminar = tk.Toplevel()
    ventana_eliminar.title("Eliminar pedido")
    ventana_eliminar.geometry("450x400")

    tk.Label(
        ventana_eliminar,
        text="ELIMINAR PEDIDO",
        font=("Arial", 18)
    ).pack(pady=15)

    tk.Label(
        ventana_eliminar,
        text="Nombre del cliente:"
    ).pack(pady=5)

    entrada_cliente = tk.Entry(
        ventana_eliminar,
        width=35
    )
    entrada_cliente.pack()

    def eliminar():

        nombre_cliente = entrada_cliente.get()
        encontrado = False

        for pedido in lista_pedidos:

            if pedido["cliente"] == nombre_cliente:

                respuesta = messagebox.askyesno(
                    "Confirmar eliminación",
                    "¿Está seguro de eliminar este pedido?\n\n"
                    "Cliente: " + pedido["cliente"] +
                    "\nProducto: " + pedido["producto"] +
                    "\nCantidad: " + str(pedido["cantidad"]) +
                    "\nTotal: Q" + str(pedido["total"])
                )

                if respuesta:

                    producto = buscar_producto(pedido["producto"])

                    if producto is not None:
                        producto["existencia"] = (
                            int(producto["existencia"]) +
                            int(pedido["cantidad"])
                        )

                    lista_pedidos.remove(pedido)

                    guardar_pedidos()
                    guardar_productos()
                    actualizar_excel()


                    messagebox.showinfo(
                        "Pedido eliminado",
                        "El pedido se eliminó correctamente"
                    )

                    ventana_eliminar.destroy()

                encontrado = True
                break

        if encontrado == False:

            messagebox.showerror(
                "Error",
                "No se encontró un pedido para ese cliente"
            )

    boton_eliminar = tk.Button(
        ventana_eliminar,
        text="Eliminar pedido",
        command=eliminar,
        width=25
    )
    boton_eliminar.pack(pady=20)

def guardar_pedido(cliente, producto, cantidad, total):
    pedido = {
        "cliente": cliente,
        "producto": producto,
        "cantidad": cantidad,
        "total": total
    }

    lista_pedidos.append(pedido)



def pedidos():
    pedido_actual={}

    ventana_pedidos = tk.Toplevel()

    ventana_pedidos.title("Pedidos")

    ventana_pedidos.geometry("500x500")


    titulo = tk.Label(
        ventana_pedidos,
        text="GESTIÓN DE PEDIDOS",
        font=("Arial", 20)
    )

    titulo.pack(pady=20)




    tk.Label(
        ventana_pedidos,
        text="Nombre del cliente:"
    ).pack(pady=5)


    entrada_cliente_pedido = tk.Entry(
        ventana_pedidos,
        width=35
    )

    entrada_cliente_pedido.pack()




    def ejecutar_busqueda():

        nombre = entrada_cliente_pedido.get()

        cliente = buscar_cliente(nombre)

        if cliente is not None:

            messagebox.showinfo(
                "Cliente encontrado",
                "Cliente: " + cliente["nombre"] +
                "\nNIT: " + cliente["nit"] +
                "\nDirección: " + cliente["direccion"]
            )

        else:

            respuesta = messagebox.askyesno(
                "Cliente no registrado",
                "El cliente no está registrado.\n\n"
                "¿Desea registrarlo ahora?"
            )

            if respuesta:

                ventana_cliente = tk.Toplevel(ventana_pedidos)
                ventana_cliente.title("Registrar cliente")
                ventana_cliente.geometry("400x350")

                tk.Label(
                    ventana_cliente,
                    text="REGISTRAR NUEVO CLIENTE",
                    font=("Arial", 16)
                ).pack(pady=15)

                tk.Label(
                    ventana_cliente,
                    text="Nombre:"
                ).pack(pady=5)

                entrada_nombre = tk.Entry(
                    ventana_cliente,
                    width=35
                )
                entrada_nombre.pack()

                entrada_nombre.insert(0, nombre)

                tk.Label(
                    ventana_cliente,
                    text="NIT:"
                ).pack(pady=5)

                entrada_nit = tk.Entry(
                    ventana_cliente,
                    width=35
                )
                entrada_nit.pack()

                tk.Label(
                    ventana_cliente,
                    text="Dirección:"
                ).pack(pady=5)

                entrada_direccion = tk.Entry(
                    ventana_cliente,
                    width=35
                )
                entrada_direccion.pack()

                def guardar_nuevo_cliente():

                    nuevo_cliente = {
                        "nombre": entrada_nombre.get(),
                        "nit": entrada_nit.get(),
                        "direccion": entrada_direccion.get()
                    }

                    lista_clientes.append(nuevo_cliente)

                    guardar_clientes()
                    actualizar_excel()

                    messagebox.showinfo(
                        "Cliente registrado",
                        "El cliente fue registrado correctamente"
                    )

                    ventana_cliente.destroy()

                    messagebox.showinfo(
                        "Continuar pedido",
                        "El cliente ya está registrado.\n"
                        "Ahora puede continuar con el pedido."
                    )    

                boton_guardar = tk.Button(
                    ventana_cliente,
                    text="Guardar cliente",
                    command=guardar_nuevo_cliente,
                    width=25
                )
                boton_guardar.pack(pady=20)



    boton_buscar_cliente = tk.Button(
        ventana_pedidos,
        text="Buscar cliente",
        command=ejecutar_busqueda,
        width=25
    )

    boton_buscar_cliente.pack(pady=15)


    tk.Label(
        ventana_pedidos,
        text="Nombre del producto:"
    ).pack(pady=5)


    entrada_producto_pedido = tk.Entry(
        ventana_pedidos,
        width=35
    )

    entrada_producto_pedido.pack()

    def ejecutar_busqueda_producto():

        nombre = entrada_producto_pedido.get()

        producto = buscar_producto(nombre)

        if producto is not None:

            messagebox.showinfo(
                "Producto encontrado",
                "Producto: " + producto["nombre"] +
                "\nPrecio: Q" + producto["precio"] +
                "\nExistencia: " + producto["existencia"]
            )

        else:

            messagebox.showerror(
                "Error",
                "No se encontró el producto"
            )


    def ejecutar_busqueda_producto():

        nombre = entrada_producto_pedido.get()

        producto = buscar_producto(nombre)

        if producto is not None:

            messagebox.showinfo(
                "Producto encontrado",
                "Producto: " + producto["nombre"] +
                "\nPrecio: Q" + producto["precio"] +
                "\nExistencia: " + producto["existencia"]
            )

        else:

            messagebox.showerror(
                "Error",
                "No se encontró el producto"
            )

    boton_buscar_producto = tk.Button(
        ventana_pedidos,
        text="Buscar producto",
        command=ejecutar_busqueda_producto,
        width=25
    )

    boton_buscar_producto.pack(pady=15)



    tk.Label(
        ventana_pedidos,
        text="Cantidad:"
    ).pack(pady=5)


    entrada_cantidad = tk.Entry(
        ventana_pedidos,
        width=35
    )

    entrada_cantidad.pack()



    def calcular_total():
        nonlocal pedido_actual 
        nombre = entrada_producto_pedido.get()
        cantidad = entrada_cantidad.get()

        producto = buscar_producto(nombre)

        if producto is None:
            messagebox.showerror(
                "Error",
                "Primero debe ingresar un producto válido"
            )
            return

        try:
            precio = float(producto["precio"])
            cantidad = int(cantidad)
            existencia = int(producto["existencia"])

            if cantidad <= 0:
                messagebox.showerror(
                    "Error",
                    "La cantidad debe ser mayor que 0"
                )
                return

            if cantidad > existencia:
                messagebox.showerror(
                    "Error",
                    "No hay suficiente existencia del producto"
                )
                return

            total = precio * cantidad
            
            pedido_actual = {
            "cliente": entrada_cliente_pedido.get(),
            "producto": producto["nombre"],
            "cantidad": cantidad,
            "total": total
        }
            
            messagebox.showinfo(
                "Total del pedido",
                "Precio: Q" + str(precio) +
                "\nCantidad: " + str(cantidad) +
                "\nTotal: Q" + str(total)
            )

        except ValueError:
            messagebox.showerror(
                "Error",
                "La cantidad debe ser un número"
            )

    boton_calcular = tk.Button(
        ventana_pedidos,
        text="Calcular total",
        command=calcular_total,
        width=25
    )

    boton_calcular.pack(pady=15)


    def ejecutar_guardar_pedido():
        if pedido_actual == {}:
            messagebox.showwarning(
                "Advertencia",
                "Primero debe calcular el total del pedido"
            )
            return

        guardar_pedido(
            pedido_actual["cliente"],
            pedido_actual["producto"],
            pedido_actual["cantidad"],
            pedido_actual["total"]
        )
        guardar_pedidos()

        producto = buscar_producto(pedido_actual["producto"])

        if producto is not None:
            producto["existencia"] = (
                int(producto["existencia"]) -
                pedido_actual["cantidad"]
            )
        guardar_productos()
        actualizar_excel()

        messagebox.showinfo(
            "Pedido guardado",
            "El pedido se guardó correctamente"
        )

    boton_guardar = tk.Button(
        ventana_pedidos,
        text="Guardar pedido",
        command=ejecutar_guardar_pedido,
        width=25
    )
    boton_guardar.pack(pady=15)

    def ver_pedidos():
        ventana_lista = tk.Toplevel()
        ventana_lista.title("Lista de pedidos")
        ventana_lista.geometry("500x500")

        titulo = tk.Label(
            ventana_lista,
            text="LISTA DE PEDIDOS",
            font=("Arial", 18)
        )
        titulo.pack(pady=15)

        if len(lista_pedidos) == 0:
            mensaje = tk.Label(
                ventana_lista,
                text="No hay pedidos registrados",
                font=("Arial", 12)
            )
            mensaje.pack(pady=20)
        else:
            for pedido in lista_pedidos:
                texto = (
                    "Cliente: " + pedido["cliente"] +
                    "\nProducto: " + pedido["producto"] +
                    "\nCantidad: " + str(pedido["cantidad"]) +
                    "\nTotal: Q" + str(pedido["total"]) +
                    "\n-----------------------------"
                )

                etiqueta = tk.Label(
                    ventana_lista,
                    text=texto,
                    font=("Arial", 11),
                    justify="left"
                )
                etiqueta.pack(pady=5)

    boton_ver_pedidos = tk.Button(
        ventana_pedidos,
        text="Ver pedidos",
        command=ver_pedidos,
        width=25
    )
    boton_ver_pedidos.pack(pady=15)

    boton_eliminar_pedido = tk.Button(
        ventana_pedidos,
        text="Eliminar pedido",
        command=eliminar_pedido,
        width=25
    )
    boton_eliminar_pedido.pack(pady=15)
def crear_cotizacion(cliente, producto, precio):

    from docx import Document

    documento = Document()

    documento.add_heading("COTIZACIÓN", level=1)

    documento.add_paragraph(
        "Estimado " + cliente
    )

    documento.add_paragraph(
        "El precio de nuestro producto " + producto
    )

    documento.add_paragraph(
        "es Q" + str(precio)
    )

    documento.save("cotización.docx")

def enviar_correo(destinatario):

    mensaje = EmailMessage()

    mensaje["From"] = CORREO_EMISOR
    mensaje["To"] = destinatario
    mensaje["Subject"] = "Cotización - Sistema de Ventas"

    mensaje.set_content(
        "Estimado cliente,\n\n"
        "Adjunto encontrará su cotización.\n\n"
        "Saludos,\n"
        "Sistema de Ventas"
    )

    with open("cotización.docx", "rb") as archivo:
        contenido = archivo.read()

    mensaje.add_attachment(
        contenido,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="cotización.docx"
    )

    try:

        servidor = smtplib.SMTP(
            "smtp.gmail.com",
            587
        )

        servidor.starttls()

        servidor.login(
            CORREO_EMISOR,
            CONTRASENA_APP
        )

        servidor.send_message(mensaje)

        servidor.quit()

        return True

    except Exception as error:

        messagebox.showerror(
            "Error al enviar",
            "No se pudo enviar la cotización.\n\n"
            "Revise el correo y la contraseña de aplicación."
        )

        return False

def enviar_respaldo(destinatario):

    mensaje = EmailMessage()

    mensaje["From"] = CORREO_EMISOR
    mensaje["To"] = destinatario
    mensaje["Subject"] = "Respaldo - Sistema de Ventas"

    mensaje.set_content(
        "Estimado coordinador,\n\n"
        "Adjunto encontrará el respaldo del sistema de ventas.\n\n"
        "Saludos,\n"
        "Sistema de Ventas"
    )

    with open("backup_sistema_ventas.xlsx", "rb") as archivo:
        contenido = archivo.read()

    mensaje.add_attachment(
        contenido,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="backup_sistema_ventas.xlsx"
    )

    try:

        servidor = smtplib.SMTP(
            "smtp.gmail.com",
            587
        )

        servidor.starttls()

        servidor.login(
            CORREO_EMISOR,
            CONTRASENA_APP
        )

        servidor.send_message(mensaje)

        servidor.quit()

        return True

    except Exception as error:

        messagebox.showerror(
            "Error al enviar",
            "No se pudo enviar el respaldo por correo."
        )

        return False

def mostrar_cotizacion():

    ventana_cotizacion = tk.Toplevel()
    ventana_cotizacion.title("Crear cotización")
    ventana_cotizacion.geometry("450x450")

    titulo = tk.Label(
        ventana_cotizacion,
        text="CREAR COTIZACIÓN",
        font=("Arial", 18)
    )
    titulo.pack(pady=15)

    tk.Label(
        ventana_cotizacion,
        text="Nombre del cliente:"
    ).pack(pady=5)

    entrada_cliente = tk.Entry(
        ventana_cotizacion,
        width=35
    )
    entrada_cliente.pack()

    tk.Label(
        ventana_cotizacion,
        text="Correo del cliente:"
    ).pack(pady=5)

    entrada_correo = tk.Entry(
        ventana_cotizacion,
        width=35
    )
    entrada_correo.pack() 

    tk.Label(
        ventana_cotizacion,
        text="Nombre del producto:"
    ).pack(pady=5)

    entrada_producto = tk.Entry(
        ventana_cotizacion,
        width=35
    )
    entrada_producto.pack()

    tk.Label(
        ventana_cotizacion,
        text="Precio del producto:"
    ).pack(pady=5)

    entrada_precio = tk.Entry(
        ventana_cotizacion,
        width=35
    )
    entrada_precio.pack()

    def buscar_precio():

        producto = buscar_producto(
            entrada_producto.get()
        )

        if producto is None:
            messagebox.showerror(
                "Error",
                "El producto no está registrado"
            )
            return

        entrada_precio.delete(0, tk.END)

        entrada_precio.insert(
            0,
            str(producto["precio"])
        )

    boton_buscar_precio = tk.Button(
        ventana_cotizacion,
        text="Buscar precio",
        command=buscar_precio,
        width=25
    )
    boton_buscar_precio.pack(pady=10)

    def generar():

        cliente = entrada_cliente.get()
        producto = entrada_producto.get()
        precio = entrada_precio.get()
        correo = entrada_correo.get()

        if cliente == "" or producto == "" or precio == "" or correo == "":
            messagebox.showwarning(
                "Advertencia",
                "Debe llenar todos los campos"
            )
            return

        crear_cotizacion(
            cliente,
            producto,
            precio
        )

        enviar_correo(correo)

        messagebox.showinfo(
            "Cotización enviada",
            "La cotización se creó y se envió correctamente."
        )

    boton_generar = tk.Button(
        ventana_cotizacion,
        text="Generar cotización",
        command=generar,
        width=25
    )
    boton_generar.pack(pady=20) 

def informes():
    ventana_informes = tk.Toplevel()
    ventana_informes.title("Informes")
    ventana_informes.geometry("500x450")

    titulo = tk.Label(
        ventana_informes,
        text="INFORMES DEL SISTEMA",
        font=("Arial", 18)
    )
    titulo.pack(pady=20)

    total_clientes = len(lista_clientes)
    total_productos = len(lista_productos)
    total_pedidos = len(lista_pedidos)

    ventas_totales = 0

    for pedido in lista_pedidos:
        ventas_totales = ventas_totales + float(pedido["total"])
        
    productos_bajos = []

    for producto in lista_productos:
        if int(producto["existencia"]) <= 5:
            productos_bajos.append(producto)

    texto_informe = (
        "Clientes registrados: " + str(total_clientes) +
        "\n\nProductos registrados: " + str(total_productos) +
        "\n\nPedidos realizados: " + str(total_pedidos) +
        "\n\nVentas totales: Q" + str(ventas_totales)
    )

    if len(productos_bajos) > 0:
        texto_informe = texto_informe + "\n\nPRODUCTOS CON POCA EXISTENCIA:\n"

        for producto in productos_bajos:
            texto_informe = (
                texto_informe +
                producto["nombre"] +
                " → " +
                str(producto["existencia"]) +
                " unidades\n"
            )
    else:
        texto_informe = (
            texto_informe +
            "\n\nNo hay productos con poca existencia."
        )
    ventas_productos = {}

    for pedido in lista_pedidos:
        producto = pedido["producto"]
        cantidad = int(pedido["cantidad"])

        if producto in ventas_productos:
            ventas_productos[producto] = (
                ventas_productos[producto] + cantidad
            )
        else:
            ventas_productos[producto] = cantidad

    if len(ventas_productos) > 0:
        producto_mas_vendido = max(
            ventas_productos,
            key=ventas_productos.get
        )

        cantidad_mas_vendida = ventas_productos[producto_mas_vendido]

        texto_informe = (
            texto_informe +
            "\n\nPRODUCTO MÁS VENDIDO:\n" +
            producto_mas_vendido +
            " → " +
            str(cantidad_mas_vendida) +
            " unidades"
        )
    else:
        texto_informe = (
            texto_informe +
            "\n\nNo hay ventas registradas."
        )
    compras_clientes = {}

    for pedido in lista_pedidos:
        cliente = pedido["cliente"]
        cantidad = int(pedido["cantidad"])

        if cliente in compras_clientes:
            compras_clientes[cliente] = (
                compras_clientes[cliente] + cantidad
            )
        else:
            compras_clientes[cliente] = cantidad

    if len(compras_clientes) > 0:
        cliente_mas_compras = max(
            compras_clientes,
            key=compras_clientes.get
        )

        cantidad_compras = compras_clientes[cliente_mas_compras]

        texto_informe = (
            texto_informe +
            "\n\nCLIENTE CON MÁS COMPRAS:\n" +
            cliente_mas_compras +
            " → " +
            str(cantidad_compras) +
            " unidades compradas"
        )
    else:
        texto_informe = (
            texto_informe +
            "\n\nNo hay compras registradas."
        )
            

    etiqueta_informe = tk.Label(
        ventana_informes,
        text=texto_informe,
        font=("Arial", 13),
        justify="left"
    )
    etiqueta_informe.pack(pady=20)

    ventas_por_cliente = {}

    for pedido in lista_pedidos:
        cliente = pedido["cliente"]
        total = float(pedido["total"])

        if cliente in ventas_por_cliente:
            ventas_por_cliente[cliente] = (
                ventas_por_cliente[cliente] + total
            )
        else:
            ventas_por_cliente[cliente] = total

    texto_ventas_cliente = "\n\nVENTAS TOTALES POR CLIENTE:\n"

    for cliente in ventas_por_cliente:
        texto_ventas_cliente = (
            texto_ventas_cliente +
            cliente +
            " → Q" +
            str(ventas_por_cliente[cliente]) +
            "\n"
        )

    ventas_por_producto = {}

    for pedido in lista_pedidos:
        producto = pedido["producto"]
        total = float(pedido["total"])

        if producto in ventas_por_producto:
            ventas_por_producto[producto] = (
                ventas_por_producto[producto] + total
            )
        else:
            ventas_por_producto[producto] = total

    texto_ventas_producto = "\n\nVENTAS TOTALES POR PRODUCTO:\n"

    for producto in ventas_por_producto:
        texto_ventas_producto = (
            texto_ventas_producto +
            producto +
            " → Q" +
            str(ventas_por_producto[producto]) +
            "\n"
        )

    etiqueta_ventas = tk.Label(
        ventana_informes,
        text=texto_ventas_cliente + texto_ventas_producto,
        font=("Arial", 12),
        justify="left"
    )
    etiqueta_ventas.pack(pady=10) 
def crear_respaldo():

    try:

        shutil.copy(
            "sistema_ventas.xlsx",
            "backup_sistema_ventas.xlsx"
        )

        messagebox.showinfo(
            "Respaldo creado",
            "El respaldo se creó correctamente.\n\n"
            "Archivo: backup_sistema_ventas.xlsx"
        )

    except FileNotFoundError:

        messagebox.showerror(
            "Error",
            "No se encontró el archivo sistema_ventas.xlsx"
        )

def varios():

    ventana_varios = tk.Toplevel()
    ventana_varios.title("Varios")
    ventana_varios.geometry("400x450")

    titulo = tk.Label(
        ventana_varios,
        text="OPCIONES VARIAS",
        font=("Arial", 18)
    )
    titulo.pack(pady=20)

    boton_respaldo = tk.Button(
        ventana_varios,
        text="Crear respaldo",
        command=crear_respaldo,
        width=25
    )
    boton_respaldo.pack(pady=20)

    tk.Label(
        ventana_varios,
        text="Correo del coordinador:"
    ).pack(pady=5)

    entrada_correo_respaldo = tk.Entry(
        ventana_varios,
        width=35
    )
    entrada_correo_respaldo.pack()

    def enviar():

        correo = entrada_correo_respaldo.get()

        if correo == "":
            messagebox.showwarning(
                "Advertencia",
                "Debe ingresar el correo del coordinador"
            )
            return

        resultado = enviar_respaldo(correo)

        if resultado:
            messagebox.showinfo(
                "Respaldo enviado",
                "El respaldo se envió correctamente."
            )

    boton_enviar_respaldo = tk.Button(
        ventana_varios,
        text="Enviar respaldo por correo",
        command=enviar,
        width=25
    )
    boton_enviar_respaldo.pack(pady=20)


ventana = tk.Tk()

ventana.title("Sistema de Ventas")

ventana.geometry("700x700")


titulo_principal = tk.Label(
    ventana,
    text="SISTEMA DE VENTAS",
    font=("Arial", 28, "bold")
)

titulo_principal.pack(pady=40)

boton_productos = tk.Button(
    ventana,
    text="Productos",
    command=productos,
    width=30
)

boton_productos.pack(pady=10)

boton_clientes = tk.Button(
    ventana,
    text="Clientes",
    command=clientes,
    width=30
)

boton_clientes.pack(pady=10)


boton_pedidos = tk.Button(
    ventana,
    text="Pedidos",
    command=pedidos,
    width=30
)

boton_pedidos.pack(pady=10)

boton_informes = tk.Button(
    ventana,
    text="Informes",
    command=informes,
    width=30
)

boton_informes.pack(pady=10)



boton_varios = tk.Button(
    ventana,
    text="Varios",
    command=varios,
    width=30
)

boton_varios.pack(pady=10)




boton_salir = tk.Button(
    ventana,
    text="Salir",
    command=ventana.destroy,
    width=30
)

boton_salir.pack(pady=20)



ventana.mainloop()
